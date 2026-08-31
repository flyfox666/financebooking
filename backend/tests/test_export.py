from io import BytesIO

from openpyxl import load_workbook


def _cell_map(ws):
    values = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0]:
            values[row[0]] = row
    return values


def test_export_trial_balance(client, auth_headers, mock_month, book):
    resp = client.get(
        "/api/reports/export",
        headers=auth_headers,
        params={"book_id": book.id, "period": "2026-08", "report": "trial-balance"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(BytesIO(resp.content)).active
    assert ws.title == "科目余额表"
    values = _cell_map(ws)
    assert values["1002"][6] == 464370.00
    assert values["2211"][6] == 0.00
    assert values["合计"][2] == 0.00
    assert values["合计"][4] == values["合计"][5]


def test_export_balance_sheet(client, auth_headers, mock_month_with_carryover, book):
    resp = client.get(
        "/api/reports/export",
        headers=auth_headers,
        params={"book_id": book.id, "period": "2026-08", "report": "balance-sheet"},
    )
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.content)).active
    assert ws.title == "资产负债表"
    values = _cell_map(ws)
    assert values["资产总计"][1] == 505586.67
    assert values["负债和所有者权益总计"][1] == 505586.67
    assert values["资产总计"][2] == 0.00


def test_export_income_statement(client, auth_headers, mock_month_with_carryover, book):
    resp = client.get(
        "/api/reports/export",
        headers=auth_headers,
        params={"book_id": book.id, "period": "2026-08", "report": "income-statement"},
    )
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.content)).active
    assert ws.title == "利润表"
    values = _cell_map(ws)
    assert values["一、营业收入"][1] == 27227.72
    assert values["四、净利润"][1] == -46133.61
    assert values["四、净利润"][2] == -46133.61


def test_export_unknown_report_rejected(client, auth_headers, mock_month, book):
    resp = client.get(
        "/api/reports/export",
        headers=auth_headers,
        params={"book_id": book.id, "period": "2026-08", "report": "nope"},
    )
    assert resp.status_code == 404
