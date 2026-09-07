from decimal import Decimal

from app.ledger import voucher_service
from app.ledger.tax import calendar as tax_calendar
from app.ledger.tax.cit import calc_cit
from app.ledger.tax.iit import calc_iit
from app.ledger.tax.stamp import calc_stamp
from app.ledger.tax.vat import calc_vat


def _post_simple_vouchers(
    db_session, book, mama_user, auditor_user, post_flow, contacts_pair
):
    customer_id = contacts_pair["customer"].id
    income = voucher_service.create_voucher(
        db_session,
        book_id=book.id,
        voucher_date="2026-08-05",
        attachment_count=1,
        lines=[
            {"summary": "开票收入", "account_code": "1122", "debit": "50000.00", "credit": "0", "contact_id": customer_id},
            {"summary": "确认收入", "account_code": "5001", "debit": "0", "credit": "49504.95"},
            {"summary": "增值税", "account_code": "2221", "debit": "0", "credit": "495.05"},
        ],
        operator_id=mama_user.id,
    )
    expense = voucher_service.create_voucher(
        db_session,
        book_id=book.id,
        voucher_date="2026-08-06",
        attachment_count=1,
        lines=[
            {"summary": "云费用", "account_code": "5401", "debit": "10000.00", "credit": "0"},
            {"summary": "付款", "account_code": "1002", "debit": "0", "credit": "10000.00"},
        ],
        operator_id=mama_user.id,
    )
    capital = voucher_service.create_voucher(
        db_session,
        book_id=book.id,
        voucher_date="2026-08-01",
        attachment_count=1,
        lines=[
            {"summary": "实缴出资", "account_code": "1002", "debit": "500000.00", "credit": "0"},
            {"summary": "实收资本", "account_code": "3001", "debit": "0", "credit": "500000.00"},
        ],
        operator_id=mama_user.id,
    )
    for voucher in (income, expense, capital):
        post_flow(voucher, mama_user, auditor_user)


def test_vat_under_threshold_exempts_general_invoices(
    client, auth_headers, db_session, book, mama_user, auditor_user, post_flow, contacts_pair
):
    from tests.mock_invoices import quarter3_under_threshold

    client.post(
        "/api/invoices/import",
        headers=auth_headers,
        params={"book_id": book.id, "kind": "sales"},
        files={"file": ("sales.xlsx", quarter3_under_threshold(), "application/vnd.ms-excel")},
    )
    _post_simple_vouchers(db_session, book, mama_user, auditor_user, post_flow, contacts_pair)

    result = client.get(
        "/api/tax/vat",
        headers=auth_headers,
        params={"book_id": book.id, "year": 2026, "quarter": 3},
    ).json()

    assert result["quarter_total_incl"] == "240000.00"
    assert result["exempt"] is True
    assert result["vat_payable"] == "495.05"
    assert result["exempt_sales_excl"] == "188118.81"
    assert result["exempt_vat"] == "1881.19"
    assert result["surtax"]["urban"] == "17.33"
    assert result["surtax"]["edu"] == "7.43"
    assert result["surtax"]["local_edu"] == "4.95"
    assert result["surtax"]["total"] == "29.71"


def test_vat_over_threshold_taxes_everything(client, auth_headers, book):
    from tests.mock_invoices import quarter3_over_threshold

    client.post(
        "/api/invoices/import",
        headers=auth_headers,
        params={"book_id": book.id, "kind": "sales"},
        files={"file": ("sales.xlsx", quarter3_over_threshold(), "application/vnd.ms-excel")},
    )
    result = client.get(
        "/api/tax/vat",
        headers=auth_headers,
        params={"book_id": book.id, "year": 2026, "quarter": 3},
    ).json()

    assert result["exempt"] is False
    assert result["vat_payable"] == "3564.36"
    assert result["surtax"]["total"] == "213.86"


def test_cit_small_micro_preferential(client, auth_headers, db_session, book, mama_user, auditor_user, post_flow, contacts_pair):
    _post_simple_vouchers(db_session, book, mama_user, auditor_user, post_flow, contacts_pair)
    result = client.get(
        "/api/tax/cit",
        headers=auth_headers,
        params={"book_id": book.id, "year": 2026, "quarter": 3, "employees": 5, "assets": 100000,
                "industry_eligible": True, "adjustments_confirmed": True},
    ).json()

    assert result["profit_ytd"] == "39504.95"
    assert result["preferential"] is True
    assert result["actual_rate"] == "5%"
    assert result["tax_total_ytd"] == "1975.25"
    assert result["prepaid_this"] == "1975.25"


def test_stamp_books_from_capital_increase(client, auth_headers, db_session, book, mama_user, auditor_user, post_flow):
    capital = voucher_service.create_voucher(
        db_session,
        book_id=book.id,
        voucher_date="2026-08-01",
        attachment_count=1,
        lines=[
            {"summary": "实缴出资", "account_code": "1002", "debit": "500000.00", "credit": "0"},
            {"summary": "实收资本", "account_code": "3001", "debit": "0", "credit": "500000.00"},
        ],
        operator_id=mama_user.id,
    )
    post_flow(capital, mama_user, auditor_user)

    result = client.post(
        "/api/tax/stamp",
        headers=auth_headers,
        params={"book_id": book.id},
        json={
            "year": 2026,
            "contracts": [{"type": "tech", "amount": "100000.00"}],
        },
    ).json()

    assert result["books_increase"] == "500000.00"
    assert result["books_tax"] == "62.50"
    assert result["contracts"][0]["tax"] == "15.00"
    assert result["total"] == "77.50"


def test_iit_cumulative_withholding(client, auth_headers):
    first = client.post(
        "/api/tax/iit",
        headers=auth_headers,
        json={"month": 1, "cumulative_income": "60000", "cumulative_deductions": "5000"},
    ).json()
    assert first["cumulative_taxable"] == "55000.00"
    assert first["tax_total_ytd"] == "2980.00"
    assert first["withhold_this_month"] == "2980.00"

    second = client.post(
        "/api/tax/iit",
        headers=auth_headers,
        json={
            "month": 2,
            "cumulative_income": "120000",
            "cumulative_deductions": "10000",
            "withheld_prev": "2980.00",
        },
    ).json()
    assert second["tax_total_ytd"] == "8480.00"
    assert second["withhold_this_month"] == "5500.00"


def test_filing_calendar_and_reminders(client, auth_headers):
    calendar = client.get("/api/tax/calendar", headers=auth_headers, params={"year": 2026}).json()
    dues = {entry["due_date"] for entry in calendar}
    assert "2026-07-15" in dues
    assert "2026-09-15" in dues
    assert "2027-01-15" in dues
    assert "2027-05-31" in dues

    reminders = client.get(
        "/api/tax/reminders", headers=auth_headers, params={"within_days": 45}
    ).json()
    assert len(reminders) > 0
    assert all("days_left" in entry for entry in reminders)


def test_tax_param_roundtrip(client, auth_headers, book):
    resp = client.post(
        "/api/tax/params",
        headers=auth_headers,
        params={"book_id": book.id},
        json={"tax": "vat", "name": "rate", "value": "0.03"},
    )
    assert resp.status_code == 200
    fetched = client.get(
        "/api/tax/params/vat/rate", headers=auth_headers, params={"book_id": book.id}
    ).json()
    assert fetched["value"] == "0.03"

    vat = client.get(
        "/api/tax/vat", headers=auth_headers, params={"book_id": book.id, "year": 2026, "quarter": 3}
    ).json()
    assert vat["rate"].startswith("3")

    client.post(
        "/api/tax/params",
        headers=auth_headers,
        params={"book_id": book.id},
        json={"tax": "vat", "name": "rate", "value": "0.01"},
    )


def test_tax_export_workbook(client, auth_headers, book):
    from tests.mock_invoices import quarter3_under_threshold

    client.post(
        "/api/invoices/import",
        headers=auth_headers,
        params={"book_id": book.id, "kind": "sales"},
        files={"file": ("sales.xlsx", quarter3_under_threshold(), "application/vnd.ms-excel")},
    )
    resp = client.get(
        "/api/tax/export",
        headers=auth_headers,
        params={"book_id": book.id, "year": 2026, "quarter": 3},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(resp.content))
    assert "增值税" in workbook.sheetnames
    assert "企业所得税" in workbook.sheetnames
    assert "申报日历" in workbook.sheetnames
    values = {row[0]: row[1] for row in workbook["增值税"].iter_rows(values_only=True) if row[0]}
    assert values["应纳增值税额"] == "495.05"

def test_calendar_2026_holidays_periods_and_year_boundary():
    from datetime import date
    entries = tax_calendar.filing_calendar(2026)
    q3 = next(r for r in entries if r["period"] == "2026Q3" and r["tax"].startswith("增值税"))
    assert q3["due_date"] == "2026-10-26"
    assert q3["deadline_status"] == "official" and q3["source_url"]
    wage = next(r for r in entries if r["period"] == "2026-08" and r["tax"].startswith("个人"))
    assert wage["due_date"] == "2026-09-15"
    january = tax_calendar.upcoming_reminders(31, date(2026, 1, 1))
    assert any(r["period"] == "2025Q4" and r["due_date"] == "2026-01-20" for r in january)
    future = tax_calendar.filing_calendar(2028)
    assert all(r["deadline_status"] == "unverified" and r["source_url"] is None for r in future)
    monthly = tax_calendar.filing_calendar(2026, vat_frequency="monthly", entity_type="individual")
    assert sum(r["tax"].startswith("增值税") for r in monthly) == 12
    assert not any("企业所得税" in r["tax"] for r in monthly)
    assert not any("营业账簿" in r["tax"] or "财务报表" in r["tax"] for r in entries)


def test_cit_missing_inputs_are_unknown(client, auth_headers, book):
    result = client.get("/api/tax/cit", headers=auth_headers,
                        params={"book_id": book.id, "year": 2026, "quarter": 3}).json()
    assert result["status"] == "pending"
    assert result["preferential"] is None
    assert result["conditions"]["employees_within_limit"] is None
    assert result["conditions"]["assets_within_limit"] is None
    assert result["prepaid_this"] is None and result["actual_rate"] is None


def test_cit_uses_pretax_profit_and_adjustment(client, auth_headers, db_session, book,
                                             mama_user, auditor_user, post_flow, contacts_pair):
    _post_simple_vouchers(db_session, book, mama_user, auditor_user, post_flow, contacts_pair)
    tax = voucher_service.create_voucher(
        db_session, book_id=book.id, voucher_date="2026-08-31", operator_id=mama_user.id, attachment_count=1,
        lines=[{"summary": "计提所得税", "account_code": "5801", "debit": "2000", "credit": "0"},
               {"summary": "应交所得税", "account_code": "2221", "debit": "0", "credit": "2000"}])
    post_flow(tax, mama_user, auditor_user)
    params = {"book_id": book.id, "year": 2026, "quarter": 3, "employees": "5.5", "assets": "100000",
              "industry_eligible": True, "adjustments_confirmed": True, "adjustment_net": "-1000", "prepaid_prev": "500"}
    response = client.get("/api/tax/cit", headers=auth_headers, params=params)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["profit_ytd"] == "39504.95"
    assert result["net_profit_ytd"] == "37504.95"
    assert result["actual_profit"] == "38504.95"
    assert result["tax_total_ytd"] == "1925.25"
    assert result["prepaid_this"] == "1425.25"
    from io import BytesIO
    from openpyxl import load_workbook
    exported = client.get("/api/tax/export", headers=auth_headers, params=params)
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content))["企业所得税"]
    values = dict(sheet.iter_rows(values_only=True))
    assert values["本期估算预缴"] == result["prepaid_this"]


def test_cit_unsupported_entity_and_policy(client, auth_headers, book, db_session):
    params = {"book_id": book.id, "year": 2028, "quarter": 3, "employees": 1, "assets": 1000,
              "industry_eligible": True, "adjustments_confirmed": True}
    result = client.get("/api/tax/cit", headers=auth_headers, params=params).json()
    assert result["status"] == "policy_unverified" and result["prepaid_this"] is None
    book.entity_type = "individual"
    db_session.commit()
    params["year"] = 2026
    result = client.get("/api/tax/cit", headers=auth_headers, params=params).json()
    assert result["status"] == "not_applicable" and result["prepaid_this"] is None


def test_cit_false_industry_and_loss(client, auth_headers, book):
    params = {"book_id": book.id, "year": 2026, "quarter": 3, "employees": 1, "assets": 1000,
              "industry_eligible": False, "adjustments_confirmed": True, "adjustment_net": "10000"}
    result = client.get("/api/tax/cit", headers=auth_headers, params=params).json()
    assert result["preferential"] is False and result["tax_total_ytd"] == "2500.00"
    params.update(industry_eligible=True, adjustment_net="-10000")
    result = client.get("/api/tax/cit", headers=auth_headers, params=params).json()
    assert result["preferential"] is True and result["prepaid_this"] == "0.00"


def test_cit_invalid_input_and_tax_book_access(client, auth_headers, book, db_session, mama_user):
    base = {"book_id": book.id, "year": 2026, "quarter": 3}
    for change in ({"quarter": 5}, {"employees": -1}, {"assets": "NaN"}, {"prepaid_prev": "Infinity"}):
        assert client.get("/api/tax/cit", headers=auth_headers, params={**base, **change}).status_code == 422
    from app.ledger.book_service import create_book
    other = create_book(db_session, name="独立测试账套", start_period="2026-08")
    token = client.post("/api/auth/login", json={"username": mama_user.username, "password": "mama123456"}).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    for endpoint in ("cit", "reminders", "calendar", "export"):
        assert client.get("/api/tax/" + endpoint, headers=headers, params={**base, "book_id": other.id}).status_code == 403


def test_tax_export_pending_is_not_zero(client, auth_headers, book):
    from io import BytesIO
    from openpyxl import load_workbook
    response = client.get("/api/tax/export", headers=auth_headers,
                          params={"book_id": book.id, "year": 2026, "quarter": 3})
    wb = load_workbook(BytesIO(response.content))
    values = dict(wb["企业所得税"].iter_rows(values_only=True))
    assert values["小微优惠"] == "待核对" and values["本期估算预缴"] == "待核对"
    assert wb["申报日历"].cell(1, 6).value == "官方来源"


def test_cit_dirty_mapping_and_general_export_blocked(client, auth_headers, book, db_session):
    from app.models.report import ReportTemplate
    params = {"book_id": book.id, "year": 2026, "quarter": 3}
    book.taxpayer_type = "general"
    db_session.commit()
    assert client.get("/api/tax/export", headers=auth_headers, params=params).status_code == 400
    book.taxpayer_type = "small_scale"
    row = db_session.query(ReportTemplate).filter_by(book_id=book.id, report="is", key="operating_revenue").one()
    row.formula = '[["9999", 1]]'
    db_session.commit()
    for endpoint in ("cit", "export"):
        result = client.get("/api/tax/" + endpoint, headers=auth_headers, params=params)
        assert result.status_code == 400 and "映射" in result.json()["detail"]
