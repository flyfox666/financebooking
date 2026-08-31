from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ledger.balances import fmt_amount
from app.ledger.exceptions import BookError
from app.models.tax import Invoice

TWO = Decimal("0.01")

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_no": ("发票号码", "票据号码"),
    "invoice_code": ("发票代码",),
    "invoice_date": ("开票日期", "发票日期", "填开日期"),
    "buyer_name": ("购买方名称", "购方名称", "购买方"),
    "buyer_tax_no": ("购买方税号", "购方税号", "购买方纳税人识别号", "统一社会信用代码/购买方纳税人识别号"),
    "seller_name": ("销售方名称", "销方名称", "销售方"),
    "seller_tax_no": ("销售方税号", "销方税号", "销售方纳税人识别号"),
    "goods_name": ("货物或应税劳务、服务名称", "货物名称", "项目名称", "品名"),
    "amount_excl": ("金额", "不含税金额"),
    "tax_rate": ("税率", "征收率"),
    "tax_amount": ("税额",),
    "amount_total": ("价税合计", "合计金额", "含税金额"),
    "invoice_type": ("发票种类", "发票类型", "票种"),
    "status": ("发票状态", "状态"),
}

HEADER_KEY = "发票号码"
MAX_FILE_SIZE = 5 * 1024 * 1024


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_decimal(value) -> Decimal | None:
    text = _clean(value).replace(",", "").replace("¥", "").replace("￥", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _normalize_rate(value) -> Decimal:
    text = _clean(value)
    if not text or text in ("免税", "不征税", "0"):
        return Decimal("0")
    text = text.replace("%", "")
    try:
        number = Decimal(text)
    except InvalidOperation:
        return Decimal("0")
    if number > Decimal("0.2"):
        number = number / Decimal("100")
    return number


def _normalize_status(value) -> str:
    text = _clean(value)
    if "红" in text:
        return "red"
    if "作废" in text:
        return "void"
    return "normal"


def _normalize_type(value) -> str:
    text = _clean(value)
    if "专" in text:
        return "special"
    if "数" in text:
        return "digital"
    return "general"


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _recompute(amount_total: Decimal, rate: Decimal) -> tuple[Decimal, Decimal]:
    if amount_total == 0:
        return Decimal("0.00"), Decimal("0.00")
    if rate <= 0:
        return amount_total.quantize(TWO, rounding=ROUND_HALF_UP), Decimal("0.00")
    excl = (amount_total / (Decimal(1) + rate)).quantize(TWO, rounding=ROUND_HALF_UP)
    return excl, amount_total - excl


def _find_header_row(rows) -> tuple[int, dict[str, int]] | None:
    for idx, row in enumerate(rows[:15]):
        cells = {_clean(cell): col for col, cell in enumerate(row)}
        if HEADER_KEY in cells:
            mapping = {}
            for field, aliases in COLUMN_ALIASES.items():
                for alias in aliases:
                    if alias in cells:
                        mapping[field] = cells[alias]
                        break
            if "invoice_no" in mapping:
                return idx, mapping
    return None


def import_invoices(
    db: Session,
    *,
    book_id: int,
    kind: str,
    content: bytes,
) -> dict:
    if kind not in ("sales", "purchase"):
        raise BookError("发票方向 kind 须为 sales 或 purchase")
    if not content:
        raise BookError("导入文件为空")
    if len(content) > MAX_FILE_SIZE:
        raise BookError("导入文件不能超过 5MB")

    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise BookError("无法解析 Excel 文件，请使用金税平台导出的原始文件")

    ws = workbook.active
    rows = list(ws.iter_rows(values_only=True))
    workbook.close()
    header = _find_header_row(rows)
    if header is None:
        raise BookError("未找到表头（需包含“发票号码”列）")
    header_idx, mapping = header

    created = updated = skipped = 0
    for row in rows[header_idx + 1:]:
        raw_no = _clean(row[mapping["invoice_no"]]) if "invoice_no" in mapping else ""
        raw_no = raw_no.replace("No.", "").replace("no.", "")
        if not raw_no or not any(_clean(cell) for cell in row):
            if not raw_no:
                continue

        invoice_no = raw_no
        existing = db.scalar(
            select(Invoice).where(
                Invoice.book_id == book_id, Invoice.kind == kind, Invoice.invoice_no == invoice_no
            )
        )

        amount_total = _to_decimal(row[mapping["amount_total"]]) if "amount_total" in mapping else None
        if amount_total is None:
            amount_excl_raw = _to_decimal(row[mapping["amount_excl"]]) if "amount_excl" in mapping else None
            tax_raw = _to_decimal(row[mapping["tax_amount"]]) if "tax_amount" in mapping else None
            if amount_excl_raw is None and tax_raw is None:
                skipped += 1
                continue
            amount_total = (amount_excl_raw or Decimal("0")) + (tax_raw or Decimal("0"))
        rate = _normalize_rate(row[mapping["tax_rate"]]) if "tax_rate" in mapping else Decimal("0")
        status = _normalize_status(row[mapping["status"]]) if "status" in mapping else "normal"
        if status == "red" and amount_total > 0:
            amount_total = -amount_total
        amount_excl, tax_amount = _recompute(amount_total, rate)
        invoice_date = (
            _parse_date(row[mapping["invoice_date"]])
            if "invoice_date" in mapping
            else None
        ) or date.today()

        fields = dict(
            invoice_type=_normalize_type(row[mapping["invoice_type"]]) if "invoice_type" in mapping else "general",
            invoice_code=_clean(row[mapping["invoice_code"]]) if "invoice_code" in mapping else None,
            invoice_date=invoice_date,
            seller_name=_clean(row[mapping["seller_name"]]) if "seller_name" in mapping else "",
            seller_tax_no=_clean(row[mapping["seller_tax_no"]]) if "seller_tax_no" in mapping else "",
            buyer_name=_clean(row[mapping["buyer_name"]]) if "buyer_name" in mapping else "",
            buyer_tax_no=_clean(row[mapping["buyer_tax_no"]]) if "buyer_tax_no" in mapping else "",
            goods_name=_clean(row[mapping["goods_name"]]) if "goods_name" in mapping else "",
            amount_total=amount_total.quantize(TWO, rounding=ROUND_HALF_UP),
            amount_excl=amount_excl,
            tax_amount=tax_amount,
            tax_rate=rate,
            status=status,
        )

        if existing is None:
            db.add(Invoice(book_id=book_id, kind=kind, invoice_no=invoice_no, **fields))
            created += 1
        else:
            for key, value in fields.items():
                setattr(existing, key, value)
            updated += 1

    db.commit()
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "kind": kind,
    }


def list_invoices(
    db: Session,
    *,
    book_id: int,
    kind: str | None = None,
    period: str | None = None,
) -> list[Invoice]:
    stmt = select(Invoice).where(Invoice.book_id == book_id)
    if kind:
        stmt = stmt.where(Invoice.kind == kind)
    if period:
        year, month = int(period[:4]), int(period[5:7])
        start = date(year, month, 1)
        end = date(year, month, calendar_month_end(year, month))
        stmt = stmt.where(Invoice.invoice_date >= start, Invoice.invoice_date <= end)
    stmt = stmt.order_by(Invoice.invoice_date, Invoice.invoice_no)
    return list(db.scalars(stmt))


def calendar_month_end(year: int, month: int) -> int:
    import calendar

    return calendar.monthrange(year, month)[1]
