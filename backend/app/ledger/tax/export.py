from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from app.ledger.balances import fmt_amount
from app.ledger.tax import calendar as tax_calendar
from app.ledger.tax.cit import calc_cit
from app.ledger.tax.iit import BRACKETS
from app.ledger.tax.stamp import calc_stamp
from app.ledger.tax.vat import calc_vat

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _save(workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _kv_sheet(workbook, title: str, pairs: list[tuple[str, str]]) -> None:
    ws = workbook.create_sheet(title)
    for key, value in pairs:
        ws.append([key, value])


def export_tax_workbook(
    db,
    *,
    book_id: int,
    year: int,
    quarter: int,
    book_name: str,
    unissued_income: Decimal = Decimal("0"),
    cit_inputs: dict | None = None,
    vat_frequency: str = "quarterly",
    entity_type: str = "company",
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "申报辅助表"
    summary.append(["LedgerAI 申报辅助表"])
    summary.append([f"账套：{book_name}", f"期间：{year}Q{quarter}"])
    summary.append(["提示：本表仅作申报核对辅助，最终申报以电子税务局为准"])

    vat = calc_vat(db, book_id=book_id, year=year, quarter=quarter, unissued_income=unissued_income)
    _kv_sheet(
        workbook,
        "增值税",
        [
            ("季度销售额（价税合计）", vat["quarter_total_incl"]),
            ("其中：专票", vat["special_total_incl"]),
            ("其中：普票", vat["general_total_incl"]),
            ("其中：未开票收入", vat["unissued_income"]),
            ("季度起征点", vat["threshold"]),
            ("是否享受免税", "是" if vat["exempt"] else "否（超起征点全额计税）"),
            ("应征增值税销售额（不含税）", vat["taxable_sales_excl"]),
            ("应纳增值税额", vat["vat_payable"]),
            ("免税销售额（不含税）", vat["exempt_sales_excl"]),
            ("免征增值税额", vat["exempt_vat"]),
        ],
    )

    surtax = vat["surtax"]
    _kv_sheet(
        workbook,
        "附加税费",
        [
            ("城建税（含减半）", surtax["urban"]),
            ("教育费附加（含减半）", surtax["edu"]),
            ("地方教育附加（含减半）", surtax["local_edu"]),
            ("合计", surtax["total"]),
        ],
    )

    cit = calc_cit(db, book_id=book_id, year=year, quarter=quarter, **(cit_inputs or {}))
    _kv_sheet(
        workbook,
        "企业所得税",
        [
            ("本年累计利润总额", cit["profit_ytd"]),
            ("核对状态", {"pending": "资料待核对", "estimated": "已核对输入的辅助估算", "not_applicable": "不适用", "policy_unverified": "政策待核验"}[cit["status"]]),
            ("预缴调整净额", cit["adjustment_net"]),
            ("计税基础", cit["taxable_base"] or "待核对"),
            ("小微优惠", "待核对" if cit["preferential"] is None else ("是" if cit["preferential"] else "否")),
            ("估算税负", cit["actual_rate"] or "待核对"),
            ("本年累计估算所得税", cit["tax_total_ytd"] or "待核对"),
            ("前期已预缴", cit["prepaid_prev"]),
            ("本期估算预缴", cit["prepaid_this"] or "待核对"),
            ("说明", "；".join(cit["hints"])),
        ],
    )

    stamp = calc_stamp(db, book_id=book_id, year=year, contracts=[])
    stamp_pairs = [(f"营业账簿（增加额 {stamp['books_increase']}）", stamp["books_tax"])]
    for row in stamp["contracts"]:
        stamp_pairs.append((f"合同 {row['type']}（金额 {row['amount']}）", row["tax"]))
    stamp_pairs.append(("印花税合计", stamp["total"]))
    _kv_sheet(workbook, "印花税", stamp_pairs)

    _kv_sheet(
        workbook,
        "个税预扣率表",
        [
            (f"全年应纳税所得额不超过 {int(limit)} 元" if limit else "超过 960,000 元部分", f"税率 {rate * 100:.0f}%  速算扣除数 {deduction}")
            for limit, rate, deduction in BRACKETS
        ],
    )

    cal = tax_calendar.filing_calendar(year, vat_frequency=vat_frequency, entity_type=entity_type)
    ws = workbook.create_sheet("申报日历")
    ws.append(["截止或预计日期", "税种", "所属期", "说明", "期限状态", "官方来源"])
    for entry in cal:
        ws.append([entry["due_date"], entry["tax"], entry["period"], entry["description"], entry["deadline_note"], entry["source_url"]])
    ws.append([None, "印花税及财务报表报送", None, "期限请按主管税务机关认定及当地通知单独核对"])

    return _save(workbook)


def fmt_param(value: Decimal) -> str:
    return fmt_amount(value)
